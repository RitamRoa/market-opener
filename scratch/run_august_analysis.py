"""
Script to run FNA pipeline for specified August 2026 dates,
fetch intraday stock prices (9:30 AM and 3:00 PM),
cross-reference with 'Ritam Stocks - Sheet2(1).csv',
and generate a styled Excel sheet in the project root.
"""

import sys
import os
import json
import csv
import logging
from datetime import datetime, timedelta

# Ensure UTF-8 output on Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT_DIR)

import yfinance as yf
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from engine.fna_pipeline import run_fna_pipeline
from engine.equity_master import resolve_entity

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("AugustRunner")

AUGUST_DATES = [
    "2026-08-01",
    "2026-08-02",
    "2026-08-03",
    "2026-08-04",
    "2026-08-07",
    "2026-08-08",
    "2026-08-09",
    "2026-08-10",
    "2026-08-11",
    "2026-08-15",
    "2026-08-16",
    "2026-08-17",
    "2026-08-18",
    "2026-08-21",
    "2026-08-22",
    "2026-08-23",
    "2026-08-24",
    "2026-08-25",
    "2026-08-28",
    "2026-08-29",
    "2026-08-30",
]

REPORTS_DIR = os.path.join(ROOT_DIR, "reports")
CSV_PATH = os.path.join(ROOT_DIR, "Ritam Stocks - Sheet2(1).csv")
OUTPUT_XLSX = os.path.join(ROOT_DIR, "august_2026_fna_stocks.xlsx")


def safe_resolve_entity(sym_or_name):
    """Safely calls resolve_entity without risking unpack TypeError."""
    try:
        res = resolve_entity(sym_or_name)
        if res and isinstance(res, (tuple, list)) and len(res) >= 2:
            return res[0], res[1]
    except Exception:
        pass
    return None, None


def load_ritam_stocks():
    """Loads tickers and company names from Ritam Stocks - Sheet2(1).csv."""
    tickers = set()
    name_to_ticker = {}
    if not os.path.exists(CSV_PATH):
        logger.warning(f"File not found: {CSV_PATH}")
        return tickers, name_to_ticker

    with open(CSV_PATH, mode="r", encoding="utf-8") as f:
        reader = csv.reader(f)
        for row in reader:
            if not row or len(row) < 1:
                continue
            sym = row[0].strip().upper()
            if sym:
                tickers.add(sym)
            if len(row) > 1:
                name = row[1].strip().lower()
                name_to_ticker[name] = sym

    logger.info(f"Loaded {len(tickers)} reference tickers from {os.path.basename(CSV_PATH)}")
    return tickers, name_to_ticker


def is_in_ritam_stocks(symbol, company_name, tickers, name_to_ticker):
    """Checks if symbol or company is in Ritam Stocks list."""
    clean_sym = symbol.replace(".NS", "").replace(".BO", "").strip().upper()
    if clean_sym in tickers:
        return "Yes"
    
    # Try resolving BSE code to NSE symbol
    if clean_sym.isdigit() or symbol.endswith(".BO"):
        res_sym, _ = safe_resolve_entity(symbol)
        if res_sym and res_sym.upper() in tickers:
            return "Yes"

    # Try name matching
    norm_name = company_name.strip().lower()
    for ref_name in name_to_ticker:
        if ref_name in norm_name or norm_name in ref_name:
            return "Yes"

    return "No"


PRICE_CACHE = {}

def get_market_prices(symbol, date_str):
    """
    Fetches 9:30 AM and 3:00 PM prices for symbol on date_str.
    Returns (price_930, price_1500, note).
    """
    cache_key = (symbol, date_str)
    if cache_key in PRICE_CACHE:
        return PRICE_CACHE[cache_key]

    clean_sym = symbol.strip()
    if clean_sym.endswith(".BO") or clean_sym.isdigit():
        resolved_sym, _ = safe_resolve_entity(clean_sym)
        if resolved_sym and not resolved_sym.isdigit():
            clean_sym = f"{resolved_sym}.NS"
        elif not clean_sym.endswith(".BO"):
            clean_sym = f"{clean_sym}.BO"
    elif not clean_sym.endswith(".NS") and not clean_sym.endswith(".BO"):
        clean_sym = f"{clean_sym}.NS"

    dt = datetime.strptime(date_str, "%Y-%m-%d")
    is_weekend = dt.weekday() in [5, 6]
    next_day = (dt + timedelta(days=1)).strftime("%Y-%m-%d")

    if is_weekend:
        day_name = dt.strftime("%A")
        res = (None, None, f"Market Closed ({day_name})")
        PRICE_CACHE[cache_key] = res
        return res

    t = yf.Ticker(clean_sym)

    # 1. Try 5m interval
    try:
        df_5m = t.history(start=date_str, end=next_day, interval="5m")
        if not df_5m.empty:
            p_930 = None
            p_1500 = None
            for idx, row in df_5m.iterrows():
                t_str = idx.strftime("%H:%M")
                if t_str in ["09:30", "09:25", "09:35"] and p_930 is None:
                    p_930 = round(float(row["Close"]), 2)
                if t_str in ["15:00", "14:55", "15:05"]:
                    p_1500 = round(float(row["Close"]), 2)
            if p_930 is not None or p_1500 is not None:
                res = (p_930, p_1500, "Traded (5m)")
                PRICE_CACHE[cache_key] = res
                return res
    except Exception:
        pass

    # 2. Try 1h interval
    try:
        df_1h = t.history(start=date_str, end=next_day, interval="1h")
        if not df_1h.empty:
            p_930 = None
            p_1500 = None
            for idx, row in df_1h.iterrows():
                t_str = idx.strftime("%H:%M")
                if t_str == "09:15":
                    p_930 = round(float(row["Open"]), 2)
                if t_str == "14:15":
                    p_1500 = round(float(row["Close"]), 2)
                elif t_str == "15:15" and p_1500 is None:
                    p_1500 = round(float(row["Open"]), 2)
            if p_930 is not None or p_1500 is not None:
                res = (p_930, p_1500, "Traded (1h)")
                PRICE_CACHE[cache_key] = res
                return res
    except Exception:
        pass

    # 3. Try daily interval
    try:
        df_1d = t.history(start=date_str, end=next_day, interval="1d")
        if df_1d.empty:
            res = (None, None, "Market Closed (Holiday)")
        else:
            c = round(float(df_1d["Close"].iloc[0]), 2)
            res = (c, c, "Traded (Daily Close)")
        PRICE_CACHE[cache_key] = res
        return res
    except Exception:
        pass

    res = (None, None, "No Price Data")
    PRICE_CACHE[cache_key] = res
    return res


def run_pipeline_for_all_dates():
    """Ensures reports exist for all 21 August dates."""
    for idx, d_str in enumerate(AUGUST_DATES, 1):
        json_file = os.path.join(REPORTS_DIR, f"fna_{d_str}.json")
        if os.path.exists(json_file):
            try:
                with open(json_file, "r", encoding="utf-8") as f:
                    data = json.load(f)
                if isinstance(data, list) and len(data) > 0:
                    logger.info(f"[{idx}/{len(AUGUST_DATES)}] Report already exists for {d_str} ({len(data)} items).")
                    continue
            except Exception:
                pass

        logger.info(f"[{idx}/{len(AUGUST_DATES)}] Running FNA Pipeline for {d_str}...")
        try:
            items, _ = run_fna_pipeline(target_date=d_str)
            logger.info(f"[{idx}/{len(AUGUST_DATES)}] Generated {len(items)} items for {d_str}.")
        except Exception as e:
            logger.error(f"FNA Pipeline error on {d_str}: {e}")


def build_excel_report():
    """Reads all reports, fetches prices, and builds the styled Excel sheet."""
    tickers, name_to_ticker = load_ritam_stocks()
    all_records = []

    for d_str in AUGUST_DATES:
        json_file = os.path.join(REPORTS_DIR, f"fna_{d_str}.json")
        items = []
        if os.path.exists(json_file):
            try:
                with open(json_file, "r", encoding="utf-8") as f:
                    items = json.load(f)
            except Exception as e:
                logger.warning(f"Failed to read {json_file}: {e}")

        logger.info(f"Compiling {len(items)} items for {d_str}...")

        for rank, it in enumerate(items, 1):
            raw_sym = (it.get("symbol") or "").strip()
            co_name = (it.get("company_name") or "").strip()
            headline = (it.get("headline") or "").strip()
            direction = (it.get("direction") or it.get("fundamental_direction") or "Positive").strip().title()

            # Clean symbol
            clean_sym = raw_sym.replace(".NS", "").replace(".BO", "")
            if clean_sym.isdigit() or raw_sym.endswith(".BO"):
                res_s, _ = safe_resolve_entity(raw_sym)
                if res_s and not res_s.isdigit():
                    clean_sym = res_s

            # Fetch intraday prices
            p930, p1500, note = get_market_prices(raw_sym, d_str)

            # Check presence in Ritam Stocks list
            in_list = is_in_ritam_stocks(raw_sym, co_name, tickers, name_to_ticker)

            all_records.append({
                "Date": d_str,
                "Rank": rank,
                "Ticker": clean_sym,
                "Full_Symbol": raw_sym,
                "Company_Name": co_name,
                "Direction": direction,
                "Price_930": p930 if p930 is not None else note,
                "Price_1500": p1500 if p1500 is not None else note,
                "Status_Note": note,
                "Present_In_Ritam_Stocks": in_list,
                "Headline": headline
            })

    # Assemble Excel
    logger.info(f"Assembling Excel spreadsheet with {len(all_records)} total rows...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "August 2026 FNA Stocks"

    headers = [
        "Date",
        "Rank",
        "Stock Ticker",
        "Company Name",
        "Positive / Negative",
        "Price @ 9:30 AM (₹)",
        "Price @ 3:00 PM (₹)",
        "Present in Ritam Stocks List",
        "Status / Trading Session",
        "FNA Headline / Key Catalyst"
    ]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Segoe UI", size=10)
    
    pos_font = Font(name="Segoe UI", size=10, bold=True, color="006100")
    pos_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    neg_font = Font(name="Segoe UI", size=10, bold=True, color="9C0006")
    neg_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    yes_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    yes_font = Font(name="Segoe UI", size=10, bold=True, color="274E13")
    no_fill = PatternFill(start_color="F3F3F3", end_color="F3F3F3", fill_type="solid")
    no_font = Font(name="Segoe UI", size=10, color="595959")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    for r_idx, rec in enumerate(all_records, start=2):
        row_vals = [
            rec["Date"],
            rec["Rank"],
            rec["Ticker"],
            rec["Company_Name"],
            rec["Direction"],
            rec["Price_930"],
            rec["Price_1500"],
            rec["Present_In_Ritam_Stocks"],
            rec["Status_Note"],
            rec["Headline"]
        ]
        ws.append(row_vals)
        ws.row_dimensions[r_idx].height = 20

        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            cell.font = cell_font
            cell.alignment = Alignment(vertical="center")

            if c_idx in [1, 2, 3, 5, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [6, 7]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '₹#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            if c_idx == 5:
                if rec["Direction"] == "Positive":
                    cell.font = pos_font
                    cell.fill = pos_fill
                elif rec["Direction"] == "Negative":
                    cell.font = neg_font
                    cell.fill = neg_fill

            if c_idx == 8:
                if rec["Present_In_Ritam_Stocks"] == "Yes":
                    cell.font = yes_font
                    cell.fill = yes_fill
                else:
                    cell.font = no_font
                    cell.fill = no_fill

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    ws.freeze_panes = "A2"
    wb.save(OUTPUT_XLSX)
    logger.info(f"Excel report saved to: {OUTPUT_XLSX}")
    print(f"\n[SUCCESS] Excel report saved to {OUTPUT_XLSX}")
    print(f"Total rows: {len(all_records)}")


def main():
    logger.info("Executing Phase 1: Pipeline execution for all August dates...")
    run_pipeline_for_all_dates()
    logger.info("Executing Phase 2: Price extraction & Excel synthesis...")
    build_excel_report()


if __name__ == "__main__":
    main()
