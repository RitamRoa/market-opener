import sys, os
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

wb_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "august_2026_fna_stocks.xlsx")
wb = openpyxl.load_workbook(wb_path)

# Check if Summary sheet already exists
if "Executive Summary" in wb.sheetnames:
    del wb["Executive Summary"]

ws_data = wb["August 2026 FNA Stocks"]
rows = list(ws_data.iter_rows(values_only=True))[1:]

# Create Executive Summary sheet
ws_sum = wb.create_sheet(title="Executive Summary", index=0)

# Styling
title_font = Font(name="Segoe UI", size=14, bold=True, color="1F4E79")
section_font = Font(name="Segoe UI", size=11, bold=True, color="1F4E79")
header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
cell_font = Font(name="Segoe UI", size=10)
bold_font = Font(name="Segoe UI", size=10, bold=True)
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

ws_sum["A1"] = "Indian Stock Market FNA — August 2026 Executive Summary"
ws_sum["A1"].font = title_font

# KPIs
total_stocks = len(rows)
yes_matches = sum(1 for r in rows if r[7] == "Yes")
pos_count = sum(1 for r in rows if r[4] == "Positive")
neg_count = sum(1 for r in rows if r[4] == "Negative")
unique_tickers = len(set(r[2] for r in rows))

kpi_data = [
    ("Total August Dates Evaluated", 21),
    ("Total Fundamental Stocks Surfaced", total_stocks),
    ("Distinct Company Tickers", unique_tickers),
    ("Matches in Ritam Stocks Watchlist", f"{yes_matches} ({yes_matches/total_stocks*100:.1f}%)"),
    ("Positive Catalysts", f"{pos_count} ({pos_count/total_stocks*100:.1f}%)"),
    ("Negative Catalysts", f"{neg_count} ({neg_count/total_stocks*100:.1f}%)")
]

ws_sum["A3"] = "Key Performance Indicators"
ws_sum["A3"].font = section_font

for i, (k, v) in enumerate(kpi_data, start=4):
    ws_sum.cell(row=i, column=1, value=k).font = bold_font
    ws_sum.cell(row=i, column=1).border = thin_border
    c2 = ws_sum.cell(row=i, column=2, value=v)
    c2.font = cell_font
    c2.border = thin_border
    c2.alignment = Alignment(horizontal="right")

# Date Breakdown
ws_sum["A12"] = "Date-by-Date Surfaced Breakdown"
ws_sum["A12"].font = section_font

breakdown_headers = ["Date", "Day of Week", "Market Status", "Stocks Surfaced", "In Ritam Watchlist", "Positive", "Negative"]
for col_idx, h in enumerate(breakdown_headers, start=1):
    c = ws_sum.cell(row=13, column=col_idx, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center")

# Group by date
from collections import defaultdict
date_groups = defaultdict(list)
for r in rows:
    date_groups[r[0]].append(r)

all_aug_dates = sorted(list(date_groups.keys()))
for r_idx, d in enumerate(all_aug_dates, start=14):
    d_rows = date_groups[d]
    dt = datetime.strptime(d, "%Y-%m-%d")
    day_name = dt.strftime("%A")
    is_weekend = dt.weekday() in [5, 6]
    status = "Market Closed" if is_weekend else "Traded Session"
    
    cnt = len(d_rows)
    yes_cnt = sum(1 for x in d_rows if x[7] == "Yes")
    pos_c = sum(1 for x in d_rows if x[4] == "Positive")
    neg_c = sum(1 for x in d_rows if x[4] == "Negative")
    
    vals = [d, day_name, status, cnt, yes_cnt, pos_c, neg_c]
    for c_idx, val in enumerate(vals, start=1):
        cell = ws_sum.cell(row=r_idx, column=c_idx, value=val)
        cell.font = cell_font
        cell.border = thin_border
        if c_idx in [1, 2, 3]:
            cell.alignment = Alignment(horizontal="center")
        else:
            cell.alignment = Alignment(horizontal="right")

for col in ws_sum.columns:
    col_letter = get_column_letter(col[0].column)
    max_len = max(len(str(cell.value or '')) for cell in col)
    ws_sum.column_dimensions[col_letter].width = max(max_len + 4, 15)

wb.save(wb_path)
print("[SUCCESS] Executive Summary tab added to workbook.")
